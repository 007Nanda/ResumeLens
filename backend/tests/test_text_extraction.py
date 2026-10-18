"""Tests for the resume text extraction service.

The PDF and DOCX fixtures build real documents with pypdf and python-docx, so
every test starts from bytes those libraries actually produced. Malformed and
encrypted cases are constructed on purpose to pin the service's error
behaviour. Nothing here touches the real upload directory.
"""

import io
import zipfile
from pathlib import Path
from typing import Callable

import pytest
from docx import Document
from pypdf.errors import DependencyError

from app.services import text_extraction
from app.services.text_extraction import (
    NoExtractableTextError,
    TextExtractionError,
    UnreadableDocumentError,
    UnsupportedFormatError,
    extract_resume_text,
    normalize_text,
)


def write_file(tmp_path: Path, name: str, data: bytes) -> Path:
    """Write ``data`` to ``tmp_path / name`` and return the resulting path."""
    path = tmp_path / name
    path.write_bytes(data)
    return path


def replace_zip_entry(data: bytes, name: str, replacement: bytes) -> bytes:
    """Rebuild a ZIP-based document with one entry replaced by other bytes."""
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        parts = {
            info.filename: source.read(info.filename)
            for info in source.infolist()
        }
    parts[name] = replacement
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for entry, payload in parts.items():
            archive.writestr(entry, payload)
    return buffer.getvalue()


# --- PDF extraction ----------------------------------------------------------


@pytest.mark.parametrize(
    "page_text", ["Alice Example", "Alice (Example)", "Alice \\ Example"]
)
def test_pdf_page_text_is_extracted(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes], page_text: str
) -> None:
    path = write_file(tmp_path, "resume.pdf", make_text_pdf_bytes((page_text,)))

    assert extract_resume_text(path) == page_text


def test_pdf_pages_are_joined_with_a_blank_line(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    data = make_text_pdf_bytes(("First Page", "Second Page"))
    path = write_file(tmp_path, "resume.pdf", data)

    assert extract_resume_text(path) == "First Page\n\nSecond Page"


def test_pdf_pages_without_text_are_skipped(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    data = make_text_pdf_bytes(("First Page", "", "Third Page"))
    path = write_file(tmp_path, "resume.pdf", data)

    assert extract_resume_text(path) == "First Page\n\nThird Page"


@pytest.mark.parametrize("pages", [(), ("",), ("", "", "")])
def test_pdf_without_any_text_raises_no_extractable_text(
    tmp_path: Path,
    make_text_pdf_bytes: Callable[..., bytes],
    pages: tuple[str, ...],
) -> None:
    path = write_file(tmp_path, "resume.pdf", make_text_pdf_bytes(pages))

    with pytest.raises(NoExtractableTextError):
        extract_resume_text(path)


def test_pdf_whitespace_only_text_is_rejected(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    data = make_text_pdf_bytes(("   ",))
    path = write_file(tmp_path, "resume.pdf", data)

    with pytest.raises(NoExtractableTextError):
        extract_resume_text(path)


# --- DOCX extraction ---------------------------------------------------------


def test_docx_paragraphs_are_extracted_in_order(
    tmp_path: Path, make_text_docx_bytes: Callable[..., bytes]
) -> None:
    data = make_text_docx_bytes(("Alice Example", "Experience"))
    path = write_file(tmp_path, "resume.docx", data)

    assert extract_resume_text(path) == "Alice Example\nExperience"


def test_docx_empty_paragraphs_keep_paragraph_breaks(
    tmp_path: Path, make_text_docx_bytes: Callable[..., bytes]
) -> None:
    data = make_text_docx_bytes(("Alice Example", "", "Experience"))
    path = write_file(tmp_path, "resume.docx", data)

    assert extract_resume_text(path) == "Alice Example\n\nExperience"


def test_docx_table_rows_are_extracted_in_document_order(
    tmp_path: Path, make_text_docx_bytes: Callable[..., bytes]
) -> None:
    data = make_text_docx_bytes(
        (
            "Alice Example",
            [("Company", "Role"), ("Acme", "Engineer")],
            "After the table",
        )
    )
    path = write_file(tmp_path, "resume.docx", data)

    assert extract_resume_text(path) == (
        "Alice Example\nCompany\tRole\nAcme\tEngineer\nAfter the table"
    )


def test_docx_without_any_text_raises_no_extractable_text(
    tmp_path: Path, make_text_docx_bytes: Callable[..., bytes]
) -> None:
    path = write_file(tmp_path, "resume.docx", make_text_docx_bytes())

    with pytest.raises(NoExtractableTextError):
        extract_resume_text(path)


def test_docx_whitespace_only_text_is_rejected(
    tmp_path: Path, make_text_docx_bytes: Callable[..., bytes]
) -> None:
    data = make_text_docx_bytes(("   ", "", "\t"))
    path = write_file(tmp_path, "resume.docx", data)

    with pytest.raises(NoExtractableTextError):
        extract_resume_text(path)


def test_short_but_readable_text_is_accepted(
    tmp_path: Path,
    make_text_pdf_bytes: Callable[..., bytes],
    make_text_docx_bytes: Callable[..., bytes],
) -> None:
    """Only entirely blank documents are rejected; short text is still text."""
    pdf = write_file(tmp_path, "resume.pdf", make_text_pdf_bytes(("Hi",)))
    docx = write_file(tmp_path, "resume.docx", make_text_docx_bytes(("Hi",)))

    assert extract_resume_text(pdf) == "Hi"
    assert extract_resume_text(docx) == "Hi"


# --- documented limits -------------------------------------------------------


def test_docx_header_and_footer_text_is_not_extracted(tmp_path: Path) -> None:
    """Headers and footers live outside the body and are a known gap."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header Marker"
    document.sections[0].footer.paragraphs[0].text = "Footer Marker"
    document.add_paragraph("Body Text")
    buffer = io.BytesIO()
    document.save(buffer)
    path = write_file(tmp_path, "resume.docx", buffer.getvalue())

    assert extract_resume_text(path) == "Body Text"


def test_docx_nested_table_text_is_not_extracted(tmp_path: Path) -> None:
    """Nested tables are a known gap; the outer cell's own text is read."""
    document = Document()
    outer_cell = document.add_table(rows=1, cols=1).cell(0, 0)
    outer_cell.paragraphs[0].text = "Outer Cell"
    nested_cell = outer_cell.add_table(rows=1, cols=1).cell(0, 0)
    nested_cell.text = "Nested Cell"
    buffer = io.BytesIO()
    document.save(buffer)
    path = write_file(tmp_path, "resume.docx", buffer.getvalue())

    text = extract_resume_text(path)
    assert "Outer Cell" in text
    assert "Nested Cell" not in text


# --- normalisation -----------------------------------------------------------


def test_normalize_text_returns_empty_for_whitespace_only_input() -> None:
    assert normalize_text("") == ""
    assert normalize_text("  \r\n\t\n\xa0 ") == ""


def test_normalize_text_normalises_line_endings_and_strips_lines() -> None:
    raw = "  Alice  \r\n Example\t\r\n"

    assert normalize_text(raw) == "Alice\nExample"


def test_normalize_text_keeps_a_single_blank_line_between_paragraphs() -> None:
    assert normalize_text("Alice\n\nExperience") == "Alice\n\nExperience"


def test_normalize_text_collapses_runs_of_blank_lines() -> None:
    assert normalize_text("Alice\n\n\n\nExperience") == "Alice\n\nExperience"
    assert normalize_text("\n\nAlice\n\n") == "Alice"


# --- unreadable documents ----------------------------------------------------


@pytest.mark.parametrize(
    "data",
    [b"", b"definitely not a resume", b"%PDF-1.4\ngarbage after the header"],
)
def test_malformed_pdf_raises_unreadable(tmp_path: Path, data: bytes) -> None:
    path = write_file(tmp_path, "broken.pdf", data)

    with pytest.raises(UnreadableDocumentError):
        extract_resume_text(path)


def test_malformed_docx_raises_unreadable(
    tmp_path: Path,
    make_docx_bytes: Callable[..., bytes],
    make_text_docx_bytes: Callable[..., bytes],
) -> None:
    real_docx = make_text_docx_bytes(("Alice Example",))
    broken_documents = {
        "garbage.docx": b"definitely not a resume",
        "truncated.docx": real_docx[: len(real_docx) // 2],
        "no-document-part.docx": make_docx_bytes(include_document=False),
        "no-relationship.docx": make_docx_bytes(),
        "corrupt-xml.docx": replace_zip_entry(
            real_docx, "word/document.xml", b"<w:document broken"
        ),
    }

    for name, data in broken_documents.items():
        path = write_file(tmp_path, name, data)
        with pytest.raises(UnreadableDocumentError):
            extract_resume_text(path)


@pytest.mark.parametrize("name", ["missing.pdf", "missing.docx"])
def test_missing_file_raises_unreadable(tmp_path: Path, name: str) -> None:
    path = tmp_path / name

    with pytest.raises(UnreadableDocumentError):
        extract_resume_text(path)


# --- unsupported formats -----------------------------------------------------


@pytest.mark.parametrize(
    "name", ["resume.txt", "resume.doc", "resume", "resume.pdf.exe"]
)
def test_unsupported_extension_raises(tmp_path: Path, name: str) -> None:
    path = write_file(tmp_path, name, b"%PDF-1.4")

    with pytest.raises(UnsupportedFormatError) as error:
        extract_resume_text(path)

    assert ".pdf" in str(error.value)
    assert ".docx" in str(error.value)


def test_uppercase_extension_is_accepted(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    data = make_text_pdf_bytes(("Alice Example",))
    path = write_file(tmp_path, "RESUME.PDF", data)

    assert extract_resume_text(path) == "Alice Example"


# --- encrypted PDFs ----------------------------------------------------------


def test_pdf_with_an_empty_user_password_is_extracted(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    """Many tools "encrypt" a PDF with an empty user password to set flags."""
    data = make_text_pdf_bytes(("Confidential Example",), password="")
    path = write_file(tmp_path, "resume.pdf", data)

    assert extract_resume_text(path) == "Confidential Example"


def test_pdf_with_a_user_password_raises_unreadable(
    tmp_path: Path, make_text_pdf_bytes: Callable[..., bytes]
) -> None:
    data = make_text_pdf_bytes(("Confidential Example",), password="s3cret")
    path = write_file(tmp_path, "resume.pdf", data)

    with pytest.raises(UnreadableDocumentError) as error:
        extract_resume_text(path)

    message = str(error.value)
    assert "password" in message
    assert "s3cret" not in message
    assert str(path) not in message


def test_pdf_encryption_that_needs_a_missing_extra_is_reported_safely(
    tmp_path: Path,
    make_text_pdf_bytes: Callable[..., bytes],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """pypdf raises DependencyError for some encrypted PDFs unless its optional
    cryptography extra is installed. That library error must not escape."""

    class ReaderWithoutCryptoSupport:
        is_encrypted = True

        def __init__(self, handle) -> None:
            self.handle = handle

        def decrypt(self, password: str) -> int:
            raise DependencyError("cryptography is required")

    monkeypatch.setattr(text_extraction, "PdfReader", ReaderWithoutCryptoSupport)
    path = write_file(tmp_path, "resume.pdf", make_text_pdf_bytes(("Alice",)))

    with pytest.raises(UnreadableDocumentError) as error:
        extract_resume_text(path)

    message = str(error.value)
    assert "encrypted" in message
    assert "cryptography" not in message
    assert str(path) not in message


# --- error messages and types ------------------------------------------------


def test_error_messages_do_not_leak_paths_or_library_details(
    tmp_path: Path,
    make_text_pdf_bytes: Callable[..., bytes],
    make_text_docx_bytes: Callable[..., bytes],
) -> None:
    failures = {
        "bad.pdf": b"%PDF-1.4\nbroken beyond repair",
        "bad.docx": b"definitely not a resume",
        "empty.pdf": make_text_pdf_bytes(("",)),
        "empty.docx": make_text_docx_bytes(),
        "resume.txt": b"plain text",
    }

    for name, data in failures.items():
        path = write_file(tmp_path, name, data)
        with pytest.raises(TextExtractionError) as error:
            extract_resume_text(path)

        message = str(error.value)
        assert message
        assert str(path) not in message
        assert "BadZipFile" not in message
        assert "PdfReadError" not in message
        assert "Traceback" not in message


def test_error_types_are_distinct() -> None:
    assert issubclass(UnsupportedFormatError, TextExtractionError)
    assert issubclass(UnreadableDocumentError, TextExtractionError)
    assert issubclass(NoExtractableTextError, TextExtractionError)
    assert not issubclass(NoExtractableTextError, UnreadableDocumentError)
    assert not issubclass(UnreadableDocumentError, NoExtractableTextError)
    assert not issubclass(UnsupportedFormatError, UnreadableDocumentError)

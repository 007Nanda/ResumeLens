"""Shared fixtures for the backend test suite.

Every fixture that touches the filesystem is rooted in pytest's ``tmp_path``,
so the suite never writes into the real ``backend/uploads`` directory and is
safe to run on a developer's machine.
"""

import io
import zipfile
from pathlib import Path
from typing import Callable, Sequence

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PageObject, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.config import DEFAULT_MAX_UPLOAD_BYTES, Settings, get_settings
from app.main import app

# Minimal byte sequences used to build upload fixtures in memory. Generating
# them here keeps binary blobs out of the repository.
_PDF_BODY = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n%%EOF\n"
_CONTENT_TYPES_XML = (
    '<?xml version="1.0"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'
)
_DOCUMENT_XML = (
    '<?xml version="1.0"?>'
    '<w:document '
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
)


@pytest.fixture
def upload_dir(tmp_path: Path) -> Path:
    """A per-test directory standing in for ``backend/uploads``."""
    return tmp_path / "uploads"


@pytest.fixture
def make_client(
    upload_dir: Path,
) -> Callable[..., TestClient]:
    """Build a TestClient whose settings are overridden, never read from env."""
    clients: list[TestClient] = []

    def _make_client(
        max_upload_bytes: int = DEFAULT_MAX_UPLOAD_BYTES,
    ) -> TestClient:
        settings = Settings(
            upload_dir=upload_dir, max_upload_bytes=max_upload_bytes
        )
        app.dependency_overrides[get_settings] = lambda: settings
        client = TestClient(app)
        clients.append(client)
        return client

    try:
        yield _make_client
    finally:
        for client in clients:
            client.close()
        app.dependency_overrides.clear()


@pytest.fixture
def client(make_client: Callable[..., TestClient]) -> TestClient:
    """A TestClient using the default settings against a temporary directory."""
    with make_client() as test_client:
        yield test_client


@pytest.fixture
def make_pdf_bytes() -> Callable[..., bytes]:
    """Build bytes that start with the required PDF signature."""

    def _make(padding: int = 0) -> bytes:
        return _PDF_BODY + (b"%" * padding)

    return _make


@pytest.fixture
def make_docx_bytes() -> Callable[..., bytes]:
    """Build a ZIP-based DOCX, optionally omitting ``word/document.xml``."""

    def _make(*, include_document: bool = True) -> bytes:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", _CONTENT_TYPES_XML)
            if include_document:
                archive.writestr("word/document.xml", _DOCUMENT_XML)
        return buffer.getvalue()

    return _make


def _add_pdf_text(writer: PdfWriter, page: PageObject, text: str) -> None:
    """Draw a single line of text near the top of a blank PDF page.

    pypdf's writer has no public API for attaching a content stream, so this
    helper uses the private ``_add_object`` call that pypdf's own test suite
    uses. Should that ever change, the fixtures fail loudly instead of quietly
    producing PDFs without text.
    """
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(content)


@pytest.fixture
def make_text_pdf_bytes() -> Callable[..., bytes]:
    """Build a real PDF that pypdf can read back, with known page text.

    ``pages`` controls the page count: each entry is the text of one page,
    and an empty entry produces a page with no text at all. Pass ``password``
    to produce an encrypted PDF; ``password=""`` reproduces the common
    "encrypted with an empty user password" shape.
    """

    def _make(
        pages: Sequence[str] = ("Alice Example",),
        *,
        password: str | None = None,
    ) -> bytes:
        writer = PdfWriter()
        for page_text in pages:
            page = writer.add_blank_page(width=612, height=792)
            if page_text:
                _add_pdf_text(writer, page, page_text)
        if password is not None:
            writer.encrypt(password)
        buffer = io.BytesIO()
        writer.write(buffer)
        return buffer.getvalue()

    return _make


@pytest.fixture
def make_text_docx_bytes() -> Callable[..., bytes]:
    """Build a real DOCX with known body content in document order.

    Each block is either a paragraph (a plain string) or a table (a non-empty
    list of rows, each row a list of cell values). Calling the fixture with no
    blocks returns an empty but valid document.
    """

    def _make(
        blocks: Sequence[str | Sequence[Sequence[str]]] = (),
    ) -> bytes:
        document = Document()
        for block in blocks:
            if isinstance(block, str):
                document.add_paragraph(block)
                continue
            rows = [list(row) for row in block]
            table = document.add_table(
                rows=len(rows), cols=max(len(row) for row in rows)
            )
            for table_row, values in zip(table.rows, rows):
                for cell, value in zip(table_row.cells, values):
                    cell.text = value
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return _make


@pytest.fixture
def stored_files(upload_dir: Path) -> Callable[[], list[Path]]:
    """List everything currently in the temporary upload directory.

    Includes any leftover ``.part`` file, so tests can assert that rejected
    uploads leave nothing behind at all.
    """

    def _stored() -> list[Path]:
        if not upload_dir.exists():
            return []
        return sorted(path for path in upload_dir.iterdir() if path.is_file())

    return _stored